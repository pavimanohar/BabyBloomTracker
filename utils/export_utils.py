"""
BabyBloom export helpers.

Exports are deliberately formatted for printing/doctor review:
- A4 portrait PDF
- BabyBloom branding/logo
- complete selected calendar dates (missing values are shown as '-')
- support for multiple, non-contiguous date ranges
"""

import os
from datetime import date, datetime, timedelta

from db import Database, SUGAR_SLOTS, VITAL_TYPES

try:
    from kivy.app import App
except ImportError:
    App = None


APP_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOGO_PATH = os.path.join(APP_DIR, "assets", "icon.png")


def get_export_dir():
    if App and App.get_running_app():
        base = os.path.join(App.get_running_app().user_data_dir, "exports")
    else:
        base = os.path.join(APP_DIR, "data", "exports")
    os.makedirs(base, exist_ok=True)
    return base


def _write_media_store_copy(local_path, mime_type):
    """Write an export to Downloads/BabyBloom and return its Android URI."""
    try:
        from kivy.utils import platform
    except ImportError:
        platform = None
    if platform != "android":
        return None

    try:
        from jnius import autoclass
        PythonActivity = autoclass("org.kivy.android.PythonActivity")
        context = PythonActivity.mActivity
        ContentValues = autoclass("android.content.ContentValues")
        MediaColumns = autoclass("android.provider.MediaStore$MediaColumns")
        Downloads = autoclass("android.provider.MediaStore$Downloads")

        values = ContentValues()
        values.put(MediaColumns.DISPLAY_NAME, os.path.basename(local_path))
        values.put(MediaColumns.MIME_TYPE, mime_type)
        values.put(MediaColumns.RELATIVE_PATH, "Download/BabyBloom")
        # Keep the MediaStore item hidden until the PDF has been copied.
        # This avoids exposing a partially-written file to another app.
        try:
            values.put(MediaColumns.IS_PENDING, 1)
        except Exception:
            pass

        resolver = context.getContentResolver()
        uri = resolver.insert(Downloads.EXTERNAL_CONTENT_URI, values)
        if uri is None:
            return None

        out_stream = None
        try:
            out_stream = resolver.openOutputStream(uri)
            if out_stream is None:
                resolver.delete(uri, None, None)
                return None
            with open(local_path, "rb") as f:
                out_stream.write(f.read())
        except Exception:
            try:
                resolver.delete(uri, None, None)
            except Exception:
                pass
            raise
        finally:
            if out_stream is not None:
                out_stream.close()

        # Publish the completed file.
        try:
            completed = ContentValues()
            completed.put(MediaColumns.IS_PENDING, 0)
            resolver.update(uri, completed, None, None)
        except Exception:
            pass

        return uri
    except Exception as e:
        print(f"MediaStore copy failed: {e}")
        return None


def save_to_public_downloads(local_path, mime_type):
    uri = _write_media_store_copy(local_path, mime_type)
    if uri is None:
        return None
    return f"Downloads/BabyBloom/{os.path.basename(local_path)}"


def share_exported_file(local_path, mime_type):
    """Open Android's real system share chooser with the PDF as a content URI."""
    try:
        from kivy.utils import platform
    except ImportError:
        platform = None

    if platform == "android":
        try:
            from jnius import autoclass, cast

            PythonActivity = autoclass("org.kivy.android.PythonActivity")
            Intent = autoclass("android.content.Intent")
            ClipData = autoclass("android.content.ClipData")
            context = PythonActivity.mActivity

            uri = _write_media_store_copy(local_path, mime_type)
            if uri is None:
                return False

            intent = Intent(Intent.ACTION_SEND)
            intent.setType(mime_type)

            # Pyjnius can select the wrong overloaded putExtra() method when a
            # raw android.net.Uri is passed directly. Explicitly cast it to
            # Parcelable so Android receives the PDF URI correctly.
            parcelable_uri = cast("android.os.Parcelable", uri)
            intent.putExtra(Intent.EXTRA_STREAM, parcelable_uri)
            intent.addFlags(Intent.FLAG_GRANT_READ_URI_PERMISSION)
            intent.setClipData(ClipData.newRawUri("BabyBloom PDF", uri))

            # The temporary URI permission is carried by the SEND intent.
            # Do not depend on queryIntentActivities()/manual grants here:
            # Android package visibility rules can make that query incomplete
            # even though the system share sheet itself can handle the intent.
            chooser = Intent.createChooser(intent, "Share BabyBloom PDF")
            chooser.addFlags(Intent.FLAG_ACTIVITY_NEW_TASK)
            context.startActivity(chooser)
            return True
        except Exception as e:
            print(f"Android share failed: {e}")
            return False

    try:
        from plyer import share
        share.share(title="BabyBloom export", filepath=local_path)
        return True
    except Exception as e:
        print(f"Desktop share failed: {e}")
        return False


def open_exported_file(local_path, mime_type):
    try:
        from kivy.utils import platform
    except ImportError:
        platform = None
    if platform != "android":
        return False

    try:
        from jnius import autoclass
        PythonActivity = autoclass("org.kivy.android.PythonActivity")
        Intent = autoclass("android.content.Intent")
        context = PythonActivity.mActivity
        uri = _write_media_store_copy(local_path, mime_type)
        if uri is None:
            return False
        intent = Intent(Intent.ACTION_VIEW)
        intent.setDataAndType(uri, mime_type)
        intent.addFlags(Intent.FLAG_GRANT_READ_URI_PERMISSION)
        intent.addFlags(Intent.FLAG_ACTIVITY_NEW_TASK)
        context.startActivity(intent)
        return True
    except Exception as e:
        print(f"open_exported_file failed: {e}")
        return False


def delete_export(local_path):
    try:
        if os.path.isfile(local_path):
            os.remove(local_path)
        return True
    except OSError as e:
        print(f"delete_export failed: {e}")
        return False


def date_range_dates(start, end):
    start_d = date.fromisoformat(start)
    end_d = date.fromisoformat(end)
    if end_d < start_d:
        raise ValueError("End date must be on or after start date.")
    cur = start_d
    while cur <= end_d:
        yield cur.isoformat()
        cur += timedelta(days=1)


def normalize_ranges(ranges):
    """Validate, de-duplicate and sort selected ranges without merging gaps."""
    normalized = []
    for start, end in ranges:
        start = date.fromisoformat(start)
        end = date.fromisoformat(end)
        if end < start:
            raise ValueError("End date must be on or after start date.")
        normalized.append((start.isoformat(), end.isoformat()))
    return sorted(set(normalized))


def range_label(start, end):
    s = date.fromisoformat(start).strftime("%d %b %Y")
    e = date.fromisoformat(end).strftime("%d %b %Y")
    return s if start == end else f"{s} – {e}"


def _rows_by_date(rows, key="log_date"):
    out = {}
    for row in rows:
        out.setdefault(row.get(key, ""), []).append(row)
    return out


def _slot_map(rows):
    out = {}
    for r in rows:
        out.setdefault(r.get("slot", ""), []).append(r)
    return out


def _format_sugar_values(rows):
    values = []
    for r in rows:
        value = r.get("value")
        if isinstance(value, (int, float)):
            values.append(f"{value:g}")
    return " / ".join(values) if values else "-"


def format_vital_value(r):
    v1, v2 = r.get("value1"), r.get("value2")
    if v2 is not None and isinstance(v1, (int, float)) and isinstance(v2, (int, float)):
        return f"{v1:g}/{v2:g}"
    if isinstance(v1, (int, float)):
        return f"{v1:g}"
    return "-"


def _vital_values(rows):
    by_type = {}
    for r in rows:
        label = VITAL_TYPES.get(r.get("vital_type", ""), {}).get("label", r.get("vital_type", ""))
        value = format_vital_value(r)
        if value != "-":
            by_type.setdefault(label, []).append(value)
    return by_type


def _generic_excel(headers, grouped_sections, sheet_title, filename_prefix, col_widths, filename=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.drawing.image import Image as XLImage

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    total_cols = len(headers)

    # Brand header.
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    ws["A1"] = "BabyBloom"
    ws["A1"].font = Font(size=20, bold=True, color="5A4A55")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    if os.path.isfile(LOGO_PATH):
        try:
            img = XLImage(LOGO_PATH)
            img.width = 42
            img.height = 42
            ws.add_image(img, "A1")
        except Exception:
            pass

    header_fill = PatternFill(start_color="F6A6C1", end_color="F6A6C1", fill_type="solid")
    range_fill = PatternFill(start_color="C9B6E4", end_color="C9B6E4", fill_type="solid")
    header_font = Font(bold=True, color="5A4A55")
    thin = Side(style="thin", color="E0C7D3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    current = 3
    for section_label, rows_data in grouped_sections:
        ws.merge_cells(start_row=current, start_column=1, end_row=current, end_column=total_cols)
        cell = ws.cell(current, 1, f"Date range: {section_label}")
        cell.fill = range_fill
        cell.font = Font(bold=True, color="5A4A55")
        cell.alignment = Alignment(horizontal="left", vertical="center")
        current += 1

        for col, header in enumerate(headers, 1):
            c = ws.cell(current, col, header)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = border
        current += 1

        for row in rows_data:
            for col, value in enumerate(row, 1):
                c = ws.cell(current, col, value)
                c.border = border
                c.alignment = Alignment(vertical="top", wrap_text=True)
            current += 1
        current += 1

    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"

    if filename is None:
        filename = f"{filename_prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    path = os.path.join(get_export_dir(), filename)
    wb.save(path)
    return path


def _pdf_branding(elements, styles, subtitle=None):
    from reportlab.platypus import Table, TableStyle, Paragraph, Spacer, Image
    from reportlab.lib import colors
    from reportlab.lib.units import mm

    if os.path.isfile(LOGO_PATH):
        try:
            logo = Image(LOGO_PATH, width=15 * mm, height=15 * mm)
            brand = Paragraph("<b>BabyBloom</b><br/><font size=8>Pregnancy Care Tracker</font>", styles["Normal"])
            t = Table([[logo, brand]], colWidths=[20 * mm, 150 * mm])
            t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
            elements.append(t)
        except Exception:
            elements.append(Paragraph("<b>BabyBloom</b>", styles["Title"]))
    else:
        elements.append(Paragraph("BabyBloom", styles["Title"]))
    if subtitle:
        elements.append(Paragraph(subtitle, styles["Normal"]))
    elements.append(Spacer(1, 7))


def _generic_pdf(headers, grouped_sections, title, filename_prefix, filename=None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, KeepTogether
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    if filename is None:
        filename = f"{filename_prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    path = os.path.join(get_export_dir(), filename)

    doc = SimpleDocTemplate(path, pagesize=A4,
                            leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm)
    styles = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=styles["BodyText"], fontSize=7, leading=9)
    elements = []
    _pdf_branding(elements, styles, title)

    for section_label, rows_data in grouped_sections:
        elements.append(Paragraph(f"<b>Date range: {section_label}</b>", styles["Heading3"]))
        elements.append(Spacer(1, 3))
        data = [[Paragraph(f"<b>{h}</b>", small) for h in headers]]
        for row in rows_data:
            data.append([Paragraph(str(v).replace("&", "&amp;"), small) for v in row])
        table = Table(data, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F6A6C1")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#5A4A55")),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E0C7D3")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FFF6EE")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        elements.append(table)
        elements.append(Spacer(1, 10))

    doc.build(elements)
    return path


def build_daily_sugar_sections(rows, ranges):
    by_date = _rows_by_date(rows)
    sections = []
    slots = list(SUGAR_SLOTS)
    headers = ["Date"] + slots
    for start, end in ranges:
        section_rows = []
        for d in date_range_dates(start, end):
            slot_map = _slot_map(by_date.get(d, []))
            section_rows.append([d] + [_format_sugar_values(slot_map.get(slot, [])) for slot in slots])
        sections.append((range_label(start, end), section_rows))
    return headers, sections


def build_daily_vitals_sections(rows, ranges):
    by_date = _rows_by_date(rows)
    headers = ["Date", "Blood Pressure", "O2", "Pulse", "Weight"]
    labels = {"Blood Pressure": "Blood Pressure", "O2": "O2", "Pulse": "Pulse", "Weight": "Weight"}
    sections = []
    for start, end in ranges:
        section_rows = []
        for d in date_range_dates(start, end):
            vals = _vital_values(by_date.get(d, []))
            section_rows.append([d] + [" / ".join(vals.get(label, [])) or "-" for label in labels])
        sections.append((range_label(start, end), section_rows))
    return headers, sections


def build_daily_generic_sections(rows, ranges, date_key="event_date", columns=None):
    by_date = _rows_by_date(rows, date_key)
    sections = []
    for start, end in ranges:
        section_rows = []
        for d in date_range_dates(start, end):
            day_rows = by_date.get(d, [])
            if not day_rows:
                section_rows.append([d] + ["-" for _ in columns])
            else:
                values = []
                for key in columns:
                    vals = [str(r.get(key, "") or "-") for r in day_rows]
                    values.append(" / ".join(vals) if vals else "-")
                section_rows.append([d] + values)
        sections.append((range_label(start, end), section_rows))
    return ["Date"] + [label for _, label in columns], sections


def export_sugar_to_excel(rows, ranges=None, filename=None):
    ranges = normalize_ranges(ranges or [(min(r["log_date"] for r in rows), max(r["log_date"] for r in rows))])
    headers, sections = build_daily_sugar_sections(rows, ranges)
    return _generic_excel(headers, sections, "Sugar Log", "sugar_log", [13] + [18] * len(SUGAR_SLOTS), filename)


def export_sugar_to_pdf(rows, ranges=None, filename=None, title="Blood Sugar Log"):
    ranges = normalize_ranges(ranges or [(min(r["log_date"] for r in rows), max(r["log_date"] for r in rows))])
    headers, sections = build_daily_sugar_sections(rows, ranges)
    return _generic_pdf(headers, sections, title, "sugar_log", filename)


def export_vitals_to_excel(rows, ranges=None, filename=None):
    ranges = normalize_ranges(ranges or [(min(r["log_date"] for r in rows), max(r["log_date"] for r in rows))])
    headers, sections = build_daily_vitals_sections(rows, ranges)
    return _generic_excel(headers, sections, "Vitals Log", "vitals_log", [13, 20, 12, 12, 12], filename)


def export_vitals_to_pdf(rows, ranges=None, filename=None, title="Vitals Log"):
    ranges = normalize_ranges(ranges or [(min(r["log_date"] for r in rows), max(r["log_date"] for r in rows))])
    headers, sections = build_daily_vitals_sections(rows, ranges)
    return _generic_pdf(headers, sections, title, "vitals_log", filename)


def export_calendar_to_excel(events, ranges=None, filename=None):
    ranges = normalize_ranges(ranges or [(min(e["event_date"] for e in events), max(e["event_date"] for e in events))])
    headers, sections = build_daily_generic_sections(events, ranges, "event_date",
                                                     [("Events", "title"), ("Type", "event_type"), ("Notes", "notes")])
    return _generic_excel(headers, sections, "Events", "events", [13, 30, 16, 45], filename)


def export_calendar_to_pdf(events, ranges=None, filename=None, title="Events"):
    ranges = normalize_ranges(ranges or [(min(e["event_date"] for e in events), max(e["event_date"] for e in events))])
    headers, sections = build_daily_generic_sections(events, ranges, "event_date",
                                                     [("Events", "title"), ("Type", "event_type"), ("Notes", "notes")])
    return _generic_pdf(headers, sections, title, "events", filename)


def export_consultation_notes_to_excel(rows, ranges=None, filename=None):
    ranges = normalize_ranges(ranges or [(min(r["log_date"] for r in rows), max(r["log_date"] for r in rows))])
    headers, sections = build_daily_generic_sections(rows, ranges, "log_date", [("Notes", "notes")])
    return _generic_excel(headers, sections, "Notes", "notes", [13, 70], filename)


def export_consultation_notes_to_pdf(rows, ranges=None, filename=None, title="Notes"):
    ranges = normalize_ranges(ranges or [(min(r["log_date"] for r in rows), max(r["log_date"] for r in rows))])
    headers, sections = build_daily_generic_sections(rows, ranges, "log_date", [("Notes", "notes")])
    return _generic_pdf(headers, sections, title, "notes", filename)



def build_combined_sections(selected_rows, ranges):
    """Build one daily table containing all selected BabyBloom categories."""
    selected = set(selected_rows)
    sugar_by_date = _rows_by_date(selected_rows.get("sugar", []), "log_date")
    vitals_by_date = _rows_by_date(selected_rows.get("vitals", []), "log_date")
    meds_by_date = _rows_by_date(selected_rows.get("medications", []), "log_date")
    events_by_date = _rows_by_date(selected_rows.get("calendar", []), "event_date")
    notes_by_date = _rows_by_date(selected_rows.get("consultation", []), "log_date")

    headers = ["Date"]
    columns = []
    if "sugar" in selected:
        headers.extend(SUGAR_SLOTS)
        columns.extend(("sugar", slot) for slot in SUGAR_SLOTS)
    if "vitals" in selected:
        for label in ["Blood Pressure", "O2", "Pulse", "Weight"]:
            headers.append(label)
            columns.append(("vitals", label))
    if "medications" in selected:
        headers.append("Medications")
        columns.append(("medications", "Medications"))
    if "calendar" in selected:
        headers.append("Events")
        columns.append(("calendar", "Events"))
    if "consultation" in selected:
        headers.append("Notes")
        columns.append(("consultation", "Notes"))

    sections = []
    for start, end in ranges:
        section_rows = []
        for d in date_range_dates(start, end):
            row = [d]
            if "sugar" in selected:
                slot_map = _slot_map(sugar_by_date.get(d, []))
                row.extend(_format_sugar_values(slot_map.get(slot, [])) for slot in SUGAR_SLOTS)
            if "vitals" in selected:
                vals = _vital_values(vitals_by_date.get(d, []))
                row.extend([" / ".join(vals.get(label, [])) or "-" for label in ["Blood Pressure", "O2", "Pulse", "Weight"]])
            if "medications" in selected:
                med_values = []
                for m in meds_by_date.get(d, []):
                    name = str(m.get("name") or "").strip()
                    if not name:
                        continue
                    dosage = f" ({m['dosage']})" if m.get("dosage") else ""
                    time = f" @ {m['scheduled_time']}" if m.get("scheduled_time") else ""
                    status = "Taken" if str(m.get("taken")) == "1" else "Not taken"
                    med_values.append(f"{name}{dosage}{time} — {status}")
                row.append(" / ".join(med_values) if med_values else "-")
            if "calendar" in selected:
                event_values = [str(e.get("title") or "").strip() for e in events_by_date.get(d, [])]
                event_values = [v for v in event_values if v]
                row.append(" / ".join(event_values) if event_values else "-")
            if "consultation" in selected:
                note_values = [str(n.get("notes") or "").strip() for n in notes_by_date.get(d, [])]
                note_values = [v for v in note_values if v]
                row.append(" / ".join(note_values) if note_values else "-")
            section_rows.append(row)
        sections.append((range_label(start, end), section_rows))
    return headers, sections


def _rows_for_report_selection(db, category, start_date, end_date):
    """Return records for one selected reading across its complete date range."""
    if category == "sugar":
        return db.get_sugar_range(start_date, end_date)
    if category == "vitals":
        return db.get_vitals_range(start_date, end_date)
    if category == "medications":
        cur = db.conn.execute(
            "SELECT * FROM medications WHERE log_date BETWEEN ? AND ? ORDER BY log_date, scheduled_time, id",
            (start_date, end_date),
        )
    elif category == "calendar":
        cur = db.conn.execute(
            "SELECT * FROM calendar_events WHERE event_date BETWEEN ? AND ? ORDER BY event_date, id",
            (start_date, end_date),
        )
    elif category == "consultation":
        cur = db.conn.execute(
            "SELECT * FROM consultation_notes WHERE log_date BETWEEN ? AND ? ORDER BY log_date, id",
            (start_date, end_date),
        )
    else:
        return []
    cols = [d[0] for d in cur.description]
    return [dict(zip(cols, row)) for row in cur.fetchall()]


def _report_section(selection, rows):
    """Create a complete date-range table. Dates without records are shown as '-'."""
    category = selection["category"]
    label = selection["label"]
    start_date = selection["start"]
    end_date = selection["end"]
    ranges = [(start_date, end_date)]

    if category == "sugar":
        headers, sections = build_daily_sugar_sections(rows, ranges)
    elif category == "vitals":
        headers, sections = build_daily_vitals_sections(rows, ranges)
    elif category == "medications":
        by_date = _rows_by_date(rows, "log_date")
        headers = ["Date", "Medication", "Dosage", "Time", "Status"]
        section_rows = []
        for d in date_range_dates(start_date, end_date):
            day_rows = by_date.get(d, [])
            if not day_rows:
                section_rows.append([d, "-", "-", "-", "-"])
                continue
            for medication in day_rows:
                name = str(medication.get("name") or "").strip() or "-"
                dosage = str(medication.get("dosage") or "").strip() or "-"
                scheduled = str(medication.get("scheduled_time") or "").strip() or "-"
                taken = medication.get("taken")
                status = "Taken" if str(taken) == "1" else "Not taken" if taken is not None else "-"
                section_rows.append([d, name, dosage, scheduled, status])
        sections = [(range_label(start_date, end_date), section_rows)]
    elif category == "calendar":
        headers, sections = build_daily_generic_sections(
            rows, ranges, "event_date",
            [("Event", "title"), ("Type", "event_type"), ("Notes", "notes")],
        )
    elif category == "consultation":
        headers, sections = build_daily_generic_sections(
            rows, ranges, "log_date", [("Notes", "notes")],
        )
    else:
        return f"{label} — {range_label(start_date, end_date)}", ["Date"], [["-"]]

    return f"{label} — {range_label(start_date, end_date)}", headers, sections[0][1]


def _export_sugar_report_pdf(rows, start_date, end_date, filename=None):
    """Create the branded, doctor-friendly Blood Sugar Log report.

    This is intentionally different from the generic export table.  Each
    glucose reading gets its own row so the meal/fasting context is preserved:
    date, reading type, value, reading time, previous meal time, fasting state,
    and notes are all retained.
    """
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.platypus import (
        SimpleDocTemplate,
        Table,
        TableStyle,
        Paragraph,
        Spacer,
        Image,
        KeepTogether,
    )
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    if not rows:
        raise ValueError("No blood sugar readings found for the selected date range.")

    # BabyBloom palette.
    pink = colors.HexColor("#F2A2BF")
    pink_dark = colors.HexColor("#C94778")
    pink_soft = colors.HexColor("#FFF1F6")
    cream = colors.HexColor("#FFFDFB")
    purple = colors.HexColor("#6C5A73")
    purple_dark = colors.HexColor("#21152E")
    green_soft = colors.HexColor("#E6F5E8")
    green_text = colors.HexColor("#2E7D46")
    border = colors.HexColor("#E9CBD8")
    muted = colors.HexColor("#766B76")
    white = colors.white

    if filename is None:
        filename = f"blood_sugar_log_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    path = os.path.join(get_export_dir(), filename)

    doc = SimpleDocTemplate(
        path,
        pagesize=A4,
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=9 * mm,
        bottomMargin=13 * mm,
        title="BabyBloom — Blood Sugar Log",
        author="BabyBloom",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "bb_sugar_title", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=19, leading=22, textColor=purple_dark, alignment=TA_CENTER,
        spaceAfter=2,
    )
    subtitle_style = ParagraphStyle(
        "bb_sugar_subtitle", parent=styles["Normal"], fontName="Helvetica",
        fontSize=9.5, leading=12, textColor=muted, alignment=TA_CENTER,
    )
    body = ParagraphStyle(
        "bb_sugar_body", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=7.5, leading=9.2, textColor=purple_dark,
    )
    body_center = ParagraphStyle(
        "bb_sugar_center", parent=body, alignment=TA_CENTER,
    )
    body_left = ParagraphStyle(
        "bb_sugar_left", parent=body, alignment=TA_LEFT,
    )
    header = ParagraphStyle(
        "bb_sugar_header", parent=body, fontName="Helvetica-Bold",
        fontSize=7.2, leading=8.4, textColor=purple_dark, alignment=TA_CENTER,
    )
    small = ParagraphStyle(
        "bb_sugar_small", parent=body, fontSize=7, leading=8.5,
        textColor=muted,
    )
    card_label = ParagraphStyle(
        "bb_card_label", parent=body, fontSize=8, leading=10,
        textColor=purple, alignment=TA_CENTER,
    )
    card_value = ParagraphStyle(
        "bb_card_value", parent=body, fontName="Helvetica-Bold",
        fontSize=14, leading=16, textColor=purple_dark, alignment=TA_CENTER,
    )

    elements = []

    # Header / branding.
    if os.path.isfile(LOGO_PATH):
        try:
            logo = Image(LOGO_PATH, width=17 * mm, height=17 * mm)
            logo_table = Table(
                [[logo, Paragraph(
                    '<font color="#C94778" size="18"><b>BabyBloom</b></font>'
                    '<br/><font color="#6C5A73" size="8">Pregnancy Care Tracker</font>',
                    body_left,
                )]],
                colWidths=[20 * mm, 95 * mm],
            )
            logo_table.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
            ]))
            elements.append(logo_table)
        except Exception:
            elements.append(Paragraph("<b>BabyBloom</b>", title_style))
    else:
        elements.append(Paragraph("<b>BabyBloom</b>", title_style))

    elements.append(Spacer(1, 2 * mm))
    elements.append(Paragraph("Blood Sugar Log", title_style))
    elements.append(Paragraph(
        "Your Health  <font color='#C94778'>♥</font>  Our Care  "
        "<font color='#C94778'>♥</font>  A Happier You",
        subtitle_style,
    ))
    elements.append(Spacer(1, 4 * mm))

    # Sort by date and reading time, preserving every individual reading.
    def _sort_key(r):
        return (str(r.get("log_date", "")), str(r.get("reading_time", "")), int(r.get("id", 0) or 0))

    rows = sorted(rows, key=_sort_key)
    numeric = [
        float(r["value"])
        for r in rows
        if isinstance(r.get("value"), (int, float))
    ]
    avg = sum(numeric) / len(numeric) if numeric else None
    lowest = min(numeric) if numeric else None
    highest = max(numeric) if numeric else None
    fasting_count = sum(1 for r in rows if bool(r.get("fasting")))
    non_fasting_count = len(rows) - fasting_count

    def fmt_num(value):
        if value is None:
            return "-"
        return f"{value:g}"

    # Information card.
    left_info = [
        [Paragraph("<b>Name</b>", body), Paragraph("Mama", body)],
        [Paragraph("<b>From</b>", body), Paragraph(date.fromisoformat(start_date).strftime("%d %b %Y"), body)],
        [Paragraph("<b>To</b>", body), Paragraph(date.fromisoformat(end_date).strftime("%d %b %Y"), body)],
    ]
    right_info = [
        [Paragraph("<b>Total Readings</b>", body), Paragraph(str(len(rows)), body)],
        [Paragraph("<b>Fasting Readings</b>", body), Paragraph(str(fasting_count), body)],
        [Paragraph("<b>Non-Fasting Readings</b>", body), Paragraph(str(non_fasting_count), body)],
    ]
    info_left = Table(left_info, colWidths=[25 * mm, 45 * mm])
    info_right = Table(right_info, colWidths=[40 * mm, 30 * mm])
    for t in (info_left, info_right):
        t.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 1),
            ("RIGHTPADDING", (0, 0), (-1, -1), 1),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ]))
    info = Table([[info_left, info_right]], colWidths=[83 * mm, 83 * mm])
    info.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), pink_soft),
        ("BOX", (0, 0), (-1, -1), 0.8, pink),
        ("ROUNDEDCORNERS", [8]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    elements.append(info)
    elements.append(Spacer(1, 5 * mm))

    # Detailed reading table.  This is the key change from the old export:
    # meal type, fasting state, reading time and previous meal time are all
    # displayed for every reading.
    table_headers = [
        "Date", "Reading", "Value<br/>(mg/dL)", "Time",
        "Prev. Meal<br/>Time", "Fasting", "Notes",
    ]
    data = [[Paragraph(h, header) for h in table_headers]]

    for r in rows:
        value = r.get("value")
        value_text = fmt_num(value) if isinstance(value, (int, float)) else "-"
        reading = str(r.get("slot") or "-")
        reading_time = str(r.get("reading_time") or "-")
        meal_time = str(r.get("previous_meal_time") or "-")
        fasting = bool(r.get("fasting"))
        fasting_text = (
            '<font color="#2E7D46"><b>Yes</b></font>'
            if fasting else
            '<font color="#C94778"><b>No</b></font>'
        )
        notes = str(r.get("notes") or "-").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        data.append([
            Paragraph(str(r.get("log_date") or "-"), body_center),
            Paragraph(reading, body_left),
            Paragraph(value_text, body_center),
            Paragraph(reading_time, body_center),
            Paragraph(meal_time, body_center),
            Paragraph(fasting_text, body_center),
            Paragraph(notes, body_left),
        ])

    col_widths = [20 * mm, 29 * mm, 17 * mm, 17 * mm, 23 * mm, 18 * mm, 62 * mm]
    detail = Table(data, colWidths=col_widths, repeatRows=1, hAlign="LEFT")
    detail_style = [
        ("BACKGROUND", (0, 0), (-1, 0), pink),
        ("TEXTCOLOR", (0, 0), (-1, 0), purple_dark),
        ("GRID", (0, 0), (-1, -1), 0.45, border),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("ALIGN", (2, 0), (5, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    for idx in range(1, len(data)):
        detail_style.append((
            "BACKGROUND", (0, idx), (-1, idx),
            white if idx % 2 else colors.HexColor("#FFF7F1"),
        ))
        fasting_value = rows[idx - 1].get("fasting")
        detail_style.append((
            "BACKGROUND", (5, idx), (5, idx),
            green_soft if bool(fasting_value) else colors.HexColor("#FDE7EE"),
        ))
    detail.setStyle(TableStyle(detail_style))
    elements.append(detail)
    elements.append(Spacer(1, 5 * mm))

    # Summary cards.
    cards = [
        ("Average", fmt_num(avg), "mg/dL", pink_soft),
        ("Lowest", fmt_num(lowest), "mg/dL", green_soft),
        ("Highest", fmt_num(highest), "mg/dL", colors.HexColor("#FDE7EE")),
        ("Total Readings", str(len(rows)), "readings", colors.HexColor("#EEE7FA")),
    ]
    card_tables = []
    for label, value, unit, bg in cards:
        card = Table([
            [Paragraph(label, card_label)],
            [Paragraph(value, card_value)],
            [Paragraph(unit, card_label)],
        ], colWidths=[40 * mm], rowHeights=[7 * mm, 9 * mm, 6 * mm])
        card.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), bg),
            ("BOX", (0, 0), (-1, -1), 0.5, border),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 2),
            ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ]))
        card_tables.append(card)

    summary = Table([card_tables], colWidths=[42 * mm] * 4)
    summary.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 1),
        ("RIGHTPADDING", (0, 0), (-1, -1), 1),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    elements.append(Paragraph("<b>Summary</b>", ParagraphStyle(
        "summary_heading", parent=styles["Heading2"], fontName="Helvetica-Bold",
        fontSize=12, leading=14, textColor=purple_dark, spaceAfter=3 * mm,
    )))
    elements.append(summary)
    elements.append(Spacer(1, 5 * mm))

    # Gentle closing section with the BabyBloom hand/baby image.
    closing_text = Table([
        [Paragraph(
            '<font color="#C94778" size="13"><b>Healthy Moms<br/>Build Brighter Tomorrows ♥</b></font>'
            '<br/><font color="#6C5A73" size="7.5">Keep this report with your medical records for discussion with your healthcare professional.</font>',
            body_left,
        ),
         Image(os.path.join(APP_DIR, "assets", "mom_baby_finger.png"), width=25 * mm, height=25 * mm)
         if os.path.isfile(os.path.join(APP_DIR, "assets", "mom_baby_finger.png")) else ""]
    ], colWidths=[130 * mm, 40 * mm])
    closing_text.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FFF8FB")),
        ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#F1C9D9")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    elements.append(closing_text)

    def draw_page(canvas, document):
        canvas.saveState()
        width, height = A4
        # Soft top accent.
        canvas.setFillColor(colors.HexColor("#FFF4F8"))
        canvas.rect(0, height - 8 * mm, width, 8 * mm, stroke=0, fill=1)
        # Footer rule and metadata.
        canvas.setStrokeColor(colors.HexColor("#E8CAD8"))
        canvas.setLineWidth(0.5)
        canvas.line(10 * mm, 9 * mm, width - 10 * mm, 9 * mm)
        canvas.setFont("Helvetica", 6.8)
        canvas.setFillColor(muted)
        canvas.drawString(10 * mm, 5.2 * mm, "BabyBloom  |  Track  ♥  Care  ♥  Grow")
        canvas.drawRightString(
            width - 10 * mm, 5.2 * mm,
            f"Generated: {datetime.now().strftime('%d %b %Y')}  |  Page {doc.page}",
        )
        canvas.restoreState()

    doc.build(elements, onFirstPage=draw_page, onLaterPages=draw_page)
    return path




def export_selected_reports_to_pdf(selections, filename=None, title="BabyBloom Report"):
    """Combine multiple reading/date-range selections into one A4 portrait PDF."""
    if not selections:
        raise ValueError("No report selections were provided.")

    # Use the dedicated branded Blood Sugar Log template when the user
    # exports the Sugar option by itself. Combined reports retain the
    # existing generic layout so multiple categories can still be merged.
    if len(selections) == 1 and selections[0].get("category") == "sugar":
        selection = selections[0]
        db = Database.instance()
        rows = _rows_for_report_selection(
            db, "sugar", selection["start"], selection["end"]
        )
        return _export_sugar_report_pdf(
            rows, selection["start"], selection["end"], filename=filename
        )

    db = Database.instance()
    grouped_sections = []
    for selection in selections:
        category = selection["category"]
        start_date = selection["start"]
        end_date = selection["end"]
        try:
            start = date.fromisoformat(start_date)
            end = date.fromisoformat(end_date)
        except ValueError as exc:
            raise ValueError("Invalid report date range.") from exc
        if start > end:
            raise ValueError("Start Date cannot be after End Date.")
        rows = _rows_for_report_selection(db, category, start_date, end_date)
        section_label, headers, section_rows = _report_section(selection, rows)
        grouped_sections.append((section_label, headers, section_rows))

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    if filename is None:
        filename = f"babybloom_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    path = os.path.join(get_export_dir(), filename)
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=10 * mm)
    styles = getSampleStyleSheet()
    small = ParagraphStyle("report_small", parent=styles["BodyText"], fontSize=8, leading=10, alignment=1)
    elements = []
    _pdf_branding(elements, styles, title)

    for section_label, headers, rows in grouped_sections:
        elements.append(Paragraph(f"<b>{section_label}</b>", styles["Heading3"]))
        elements.append(Spacer(1, 3))
        data = [[Paragraph(f"<b>{h}</b>", small) for h in headers]]
        for row in rows:
            data.append([Paragraph(str(v).replace("&", "&amp;"), small) for v in row])
        table = Table(data, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F6A6C1")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#5A4A55")),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E0C7D3")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FFF6EE")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        elements.append(table)
        elements.append(Spacer(1, 10))

    doc.build(elements)
    return path

def export_combined_to_excel(selected_rows, ranges=None, filename=None):
    date_values = []
    date_keys = {"sugar": "log_date", "vitals": "log_date", "medications": "log_date",
                 "calendar": "event_date", "consultation": "log_date"}
    for key, rows in selected_rows.items():
        date_values.extend(r.get(date_keys[key]) for r in rows if r.get(date_keys[key]))
    if not date_values:
        raise ValueError("No records available for the selected categories.")
    ranges = normalize_ranges(ranges or [(min(date_values), max(date_values))])
    headers, sections = build_combined_sections(selected_rows, ranges)
    widths = [13] + [20] * (len(headers) - 1)
    return _generic_excel(headers, sections, "BabyBloom Report", "babybloom_report", widths, filename)


def export_combined_to_pdf(selected_rows, ranges=None, filename=None, title="BabyBloom Report"):
    date_values = []
    date_keys = {"sugar": "log_date", "vitals": "log_date", "medications": "log_date",
                 "calendar": "event_date", "consultation": "log_date"}
    for key, rows in selected_rows.items():
        date_values.extend(r.get(date_keys[key]) for r in rows if r.get(date_keys[key]))
    if not date_values:
        raise ValueError("No records available for the selected categories.")
    ranges = normalize_ranges(ranges or [(min(date_values), max(date_values))])
    headers, sections = build_combined_sections(selected_rows, ranges)
    return _generic_pdf(headers, sections, title, "babybloom_report", filename)

def export_medications_to_excel(rows, ranges=None, filename=None):
    ranges = normalize_ranges(ranges or [(min(r["log_date"] for r in rows), max(r["log_date"] for r in rows))])
    # Medication export is useful from the Calendar screen even though it has no standalone export tab.
    headers, sections = build_daily_generic_sections(rows, ranges, "log_date",
                                                     [("Medication", "name"), ("Dosage", "dosage"),
                                                      ("Time", "scheduled_time"), ("Taken", "taken")])
    sections = [(label, [[r[0], r[1], r[2], r[3], "Yes" if str(r[4]) == "1" else r[4]] for r in data])
                for label, data in sections]
    return _generic_excel(["Date", "Medication", "Dosage", "Time", "Taken"], sections,
                          "Medications", "medications", [13, 24, 18, 14, 10], filename)
